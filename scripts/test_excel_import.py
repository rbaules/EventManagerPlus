from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook

from services.authorization_service import capacidades_rol
from services.excel_import_service import calcular_hash_payload, construir_payload_importacion, ejecutar_importacion, generar_archivo_errores, leer_archivo_excel, preview_coincide_contexto, validar_contexto_importacion, vincular_preview_contexto
from services.excel_template_service import CANONICAL_HEADERS, IMPORT_SHEET, generar_plantilla_excel
from views.excel_import_view import excel_import_view


def workbook_bytes(rows=(), headers=CANONICAL_HEADERS, sheet_name=IMPORT_SHEET) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    sheet.append(list(headers))
    for row in rows: sheet.append(list(row))
    output = BytesIO(); book.save(output); book.close()
    return output.getvalue()


def row(code="001", recipient="Familia Pérez", order=1, name="Juan Pérez", phone="001234", email="juan@example.com", primary="Sí", table="01", table_name="Mesa Uno"):
    return (code, recipient, order, name, phone, email, primary, table, table_name)


def context(role="Administrador", phase="Pre_evento", event_state="Activo", user_state="Activo", account_state="Activo"):
    event = {"cuenta_id": 1, "evento_id": 2, "nombre_evento": "Boda", "fase_evento": phase, "estado": event_state, "rol": role}
    return {"usr_usuario_id": "u1", "usr_estado": user_state, "rol_global_calculado": role, "cuenta_actual": {"cuenta_id": 1, "nombre_cuenta": "Cuenta", "estado": account_state}, "evento_actual": event, "eventos_permitidos": [event]}


def test_template() -> None:
    data = generar_plantilla_excel()
    book = load_workbook(BytesIO(data))
    assert book.sheetnames == ["Importación", "Instrucciones"]
    sheet = book["Importación"]
    assert tuple(c.value for c in sheet[1]) == CANONICAL_HEADERS
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:I1"
    assert len(sheet.data_validations.dataValidation) == 1
    assert sheet["A2"].number_format == sheet["E2"].number_format == sheet["H2"].number_format == "@"
    assert "exactamente un principal" in book["Instrucciones"]["B5"].value
    book.close()


def test_files_and_normalization() -> None:
    valid = leer_archivo_excel("valid.xlsx", workbook_bytes([row(), row("002", "Oficina", 1, "Ana Ruiz", "0009", "ANA@EXAMPLE.COM", "Si", None, None)]))
    assert valid.is_valid and valid.summary.invitations == 2 and valid.summary.guests == 2 and valid.summary.tables == 1
    assert valid.rows[0].invitation_code == "001" and valid.rows[0].phone == "001234" and valid.rows[0].table_code == "01"
    assert valid.rows[1].email == "ana@example.com" and valid.rows[1].is_primary is True
    assert leer_archivo_excel("bad.xls", b"bad").errors
    assert leer_archivo_excel("bad.xlsx", b"bad").errors
    assert leer_archivo_excel("empty.xlsx", b"").errors
    assert leer_archivo_excel("large.xlsx", workbook_bytes([row()]), max_size=10).errors[0].code == "tamano_excedido"
    assert leer_archivo_excel("missing.xlsx", workbook_bytes([], sheet_name="Otra")).errors[0].code == "hoja_faltante"
    assert leer_archivo_excel("nodata.xlsx", workbook_bytes()).errors
    assert leer_archivo_excel("limit.xlsx", workbook_bytes([row(), row("002", "Dos", 1, "Dos")]), max_rows=1).errors
    reordered = tuple(reversed(CANONICAL_HEADERS))
    values = dict(zip(CANONICAL_HEADERS, row()))
    assert leer_archivo_excel("reordered.xlsx", workbook_bytes([[values[h] for h in reordered]], reordered)).is_valid
    assert leer_archivo_excel("missingcol.xlsx", workbook_bytes([row()[:-1]], CANONICAL_HEADERS[:-1])).errors
    assert leer_archivo_excel("extra.xlsx", workbook_bytes([(*row(), "x")], (*CANONICAL_HEADERS, "Extra"))).errors
    duplicate_headers = (*CANONICAL_HEADERS[:-1], CANONICAL_HEADERS[0])
    assert any(e.code == "encabezado_duplicado" for e in leer_archivo_excel("duphead.xlsx", workbook_bytes([row()], duplicate_headers)).errors)


def test_rules() -> None:
    invalid = [
        row(order=1.5), row(code="002", email="bad"), row(code="003", primary="X"),
        row(code="004", table="01", table_name=None), row(code="005", table=None, table_name="Mesa"),
        row(code="006", name=""), row(code="007", table_name="X" * 31),
    ]
    codes = {e.code for e in leer_archivo_excel("invalid.xlsx", workbook_bytes(invalid)).errors}
    assert {"orden_invalido", "email_invalido", "principal_invalido", "nombre_mesa_requerido", "mesa_id_requerido", "requerido_o_tipo", "longitud_excedida"} <= codes
    no_primary = leer_archivo_excel("none.xlsx", workbook_bytes([row(primary="No")]))
    assert any(e.code == "cantidad_principales" for e in no_primary.errors)
    two_primary = leer_archivo_excel("two.xlsx", workbook_bytes([row(), row(order=2, name="Ana")]))
    assert any(e.code == "cantidad_principales" for e in two_primary.errors)
    duplicate_order = leer_archivo_excel("order.xlsx", workbook_bytes([row(), row(order=1, name="Ana", primary="No")]))
    assert any(e.code == "orden_duplicado" for e in duplicate_order.errors)
    contradictory = leer_archivo_excel("recipient.xlsx", workbook_bytes([row(), row(recipient="Otra", order=2, name="Ana", primary="No")]))
    assert any(e.code == "destinatario_contradictorio" for e in contradictory.errors)
    duplicate_name = leer_archivo_excel("names.xlsx", workbook_bytes([row(), row(code="002", recipient="Dos", name="  JUAN   PEREZ  ")]))
    assert any(e.code == "nombre_duplicado" for e in duplicate_name.errors)
    differentiated = leer_archivo_excel("different.xlsx", workbook_bytes([row(name="Juan Pérez - Familia"), row(code="002", recipient="Dos", name="Juan Pérez - Oficina")]))
    assert differentiated.is_valid
    inconsistent_table = leer_archivo_excel("tables.xlsx", workbook_bytes([row(), row(code="002", recipient="Dos", name="Ana", table_name="Otra")]))
    assert any(e.code == "mesa_inconsistente" for e in inconsistent_table.errors)
    same_name = leer_archivo_excel("table-names.xlsx", workbook_bytes([row(), row(code="002", recipient="Dos", name="Ana", table="02")]))
    assert any(e.code == "nombre_mesa_duplicado" for e in same_name.errors)
    assert leer_archivo_excel("no-table.xlsx", workbook_bytes([row(table=None, table_name=None)])).is_valid


def test_context_security_and_ui() -> None:
    assert capacidades_rol("Master").puede_validar_archivo_importacion
    assert capacidades_rol("Administrador").puede_descargar_plantilla_importacion
    assert not capacidades_rol("Operador").puede_ver_importacion_excel
    assert not capacidades_rol("Consulta").puede_validar_archivo_importacion
    assert capacidades_rol("Master").puede_ejecutar_importacion
    assert validar_contexto_importacion(context("Master"))[0]
    assert validar_contexto_importacion(context("Administrador"))[0]
    for role in ("Operador", "Consulta"): assert not validar_contexto_importacion(context(role))[0]
    for phase in ("En_proceso", "Post_evento"): assert not validar_contexto_importacion(context(phase=phase))[0]
    assert not validar_contexto_importacion(context(event_state="Inactivo"))[0]
    assert not validar_contexto_importacion(context(), full_mode=False)[0]
    preview = leer_archivo_excel("valid.xlsx", workbook_bytes([row()]))
    control = excel_import_view(context(), "valid", "", "valid.xlsx", 100, preview, lambda: None, lambda: None, lambda: None)
    assert control is not None
    csv_data = generar_archivo_errores(leer_archivo_excel("inject.xlsx", workbook_bytes([row(primary="=1+1")]))).decode("utf-8-sig")
    assert "'=1+1" in csv_data
    assert leer_archivo_excel("../escape.xlsx", workbook_bytes([row()])).errors[0].code == "extension_invalida"


class RpcResponse:
    def __init__(self, data): self.data = data


class RpcCall:
    def __init__(self, owner): self.owner = owner
    def execute(self):
        if self.owner.error: raise self.owner.error
        return RpcResponse({"ok": True, "mesas_creadas": 1, "invitaciones_creadas": 1, "invitados_creados": 1})


class FakeRpc:
    def __init__(self, error=None): self.calls = []; self.error = error
    def rpc(self, name, params): self.calls.append((name, params)); return RpcCall(self)


def test_transaction_contract() -> None:
    ctx = context("Administrador")
    preview = vincular_preview_contexto(leer_archivo_excel("valid.xlsx", workbook_bytes([row()])), ctx)
    assert preview_coincide_contexto(preview, ctx)
    payload = construir_payload_importacion(preview)
    assert payload["version"] == 1 and payload["mesas"][0]["codigo_externo"] == "01"
    guest = payload["invitaciones"][0]["invitados"][0]
    assert guest["orden"] == 1 and "ivt_invitado_id" not in guest and "cuenta_id" not in payload
    assert len(calcular_hash_payload(payload)) == 64
    fake = FakeRpc()
    result = ejecutar_importacion(fake, ctx, preview)
    assert result.ok and result.guests_created == 1 and len(fake.calls) == 1
    assert fake.calls[0][0] == "evp_importar_evento_desde_json"
    changed = context("Administrador"); changed["evento_actual"]["evento_id"] = 99
    assert ejecutar_importacion(FakeRpc(), changed, preview).code == "CONTEXT_CHANGED"
    assert ejecutar_importacion(FakeRpc(), context("Operador"), preview).code == "IMPORT_FORBIDDEN"
    class ApiError(Exception):
        code = "P0001"; message = "EVENT_NOT_EMPTY"; details = ""
    assert ejecutar_importacion(FakeRpc(ApiError()), ctx, preview).message == "Este evento ya contiene información y no admite importación inicial."
    migration = (ROOT / "supabase" / "migrations" / "202608040001_excel_import_rpc.sql").read_text(encoding="utf-8")
    assert "SECURITY DEFINER" in migration and "auth.uid()" in migration
    assert "pg_try_advisory_xact_lock" in migration and "FOR UPDATE" in migration
    assert "REVOKE ALL" in migration and "GRANT EXECUTE" in migration
    assert "SET search_path = pg_catalog\n" in migration
    assert "pg_catalog, public" not in migration
    assert "NOT coalesce(v_usuario.usr_es_usuario_master, false)" in migration
    assert "v_normalized IS NULL OR v_normalized = ''" in migration
    assert "IMPORT_INTERNAL_ERROR" in migration
    assert "v_constraint_name = 'ux_evp_ivt_nombre_evento_activo'" in migration
    service_source = (ROOT / "services" / "excel_import_service.py").read_text(encoding="utf-8")
    assert ".rpc(RPC_IMPORT_NAME" in service_source
    assert ".insert(" not in service_source and ".update(" not in service_source and ".delete(" not in service_source


def main() -> None:
    test_template(); test_files_and_normalization(); test_rules(); test_context_security_and_ui(); test_transaction_contract()
    print("OK - Excel import: dependency, template, files, normalization, groups, roles, UI and security.")


if __name__ == "__main__":
    main()
