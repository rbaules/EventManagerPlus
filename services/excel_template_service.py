from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


IMPORT_SHEET = "Importación"
INSTRUCTIONS_SHEET = "Instrucciones"
TEMPLATE_FILENAME = "EventPlus_Plantilla_Importacion.xlsx"
CANONICAL_HEADERS = (
    "Código de invitación",
    "Destinatario de la invitación",
    "Orden del invitado",
    "Nombre del invitado",
    "Teléfono del invitado",
    "Email del invitado",
    "Es invitado principal?",
    "Mesa ID",
    "Nombre de mesa",
)
REQUIRED_HEADERS = frozenset((CANONICAL_HEADERS[0], CANONICAL_HEADERS[1], CANONICAL_HEADERS[2], CANONICAL_HEADERS[3], CANONICAL_HEADERS[6]))


def generar_plantilla_excel() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = IMPORT_SHEET
    header_fill = PatternFill("solid", fgColor="17365D")
    required_fill = PatternFill("solid", fgColor="D9EAF7")
    optional_fill = PatternFill("solid", fgColor="F2F2F2")
    for column, header in enumerate(CANONICAL_HEADERS, 1):
        cell = sheet.cell(1, column, header)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for row in range(2, 5002):
            sheet.cell(row, column).fill = required_fill if header in REQUIRED_HEADERS else optional_fill
    widths = (24, 32, 20, 30, 23, 30, 24, 16, 24)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:I1"
    for col in ("A", "E", "H"):
        for row in range(2, 5002):
            sheet[f"{col}{row}"].number_format = "@"
    for row in range(2, 5002):
        sheet[f"C{row}"].number_format = "0"
    validation = DataValidation(type="list", formula1='"Sí,No"', allow_blank=False)
    validation.error = "Seleccione Sí o No."
    validation.errorTitle = "Valor no permitido"
    validation.prompt = "Indique si es el invitado principal."
    validation.promptTitle = "Principal"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add("G2:G5001")

    instructions = workbook.create_sheet(INSTRUCTIONS_SHEET)
    lines = (
        ("Plantilla de importación EventPlus", "Complete una fila por invitado. No agregue, elimine ni renombre columnas."),
        ("Uso permitido", "Solo evento Activo en fase Pre_evento, inicialmente vacío, en modo FULL."),
        ("Obligatorias", "Código, Destinatario, Orden, Nombre y Es invitado principal?."),
        ("Opcionales", "Teléfono, Email, Mesa ID y Nombre de mesa (el nombre es obligatorio si existe Mesa ID)."),
        ("Principal", "Use Sí o No. Debe existir exactamente un principal por invitación."),
        ("Mesas", "Un Mesa ID siempre usa el mismo nombre; puede dejar ambos vacíos para invitados sin mesa."),
        ("Texto", "Capture teléfono, código y Mesa ID como texto para conservar ceros iniciales."),
        ("Máximos", "Código 50; destinatario 100; invitado 80; teléfono 20; email 254; Mesa ID 50; mesa 30 caracteres."),
        ("Política", "Se valida el archivo completo antes de importar y no se mezcla con datos existentes."),
        ("Seguridad", "No use fórmulas, macros ni enlaces externos. En 7B no se escriben datos."),
    )
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 110
    for row, (title, detail) in enumerate(lines, 1):
        instructions.cell(row, 1, title).font = Font(bold=True, color="17365D")
        instructions.cell(row, 2, detail).alignment = Alignment(wrap_text=True, vertical="top")
    instructions.freeze_panes = "A2"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
