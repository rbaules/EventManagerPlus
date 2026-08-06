from __future__ import annotations

import csv
import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from io import BytesIO, StringIO
from pathlib import Path
from dataclasses import replace
from datetime import datetime, timezone
from typing import Any
from zipfile import BadZipFile, is_zipfile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from models.excel_import_models import (
    ImportExecutionResult, ImportPreview, ImportRowNormalized, ImportRowRaw, ImportSummary,
    ImportValidationError, ImportValidationWarning, InvitationImportGroup,
    TableImportGroup,
)
from services.authorization_service import evento_autorizado, puede_ejecutar_importacion, puede_validar_archivo_importacion
from services.excel_template_service import CANONICAL_HEADERS, IMPORT_SHEET


MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_DATA_ROWS = 5000
PREVIEW_ROW_LIMIT = 100
RPC_IMPORT_NAME = "evp_importar_evento_desde_json"
_EMAIL = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
_LIMITS = {
    CANONICAL_HEADERS[0]: 50, CANONICAL_HEADERS[1]: 100,
    CANONICAL_HEADERS[3]: 80, CANONICAL_HEADERS[4]: 20,
    CANONICAL_HEADERS[5]: 254, CANONICAL_HEADERS[7]: 50,
    CANONICAL_HEADERS[8]: 30,
}


def normalizar_texto(value: Any) -> str:
    text = " ".join(str(value or "").strip().split()).lower()
    return "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))


def _display(value: Any) -> str:
    value = "" if value is None else str(value)
    value = value.replace("\r", " ").replace("\n", " ").replace("\x00", "")
    return value[:200]


def validar_contexto_importacion(contexto: dict[str, Any] | None, *, full_mode: bool = True) -> tuple[bool, str]:
    if not full_mode:
        return False, "La importación Excel solo está disponible en modo FULL."
    if not contexto or not contexto.get("usr_usuario_id"):
        return False, "Se requiere un usuario autenticado y activo."
    if contexto.get("usr_estado") != "Activo":
        return False, "El usuario debe estar Activo."
    cuenta = contexto.get("cuenta_actual") or {}
    if cuenta.get("estado") != "Activo":
        return False, "La cuenta activa no es válida."
    if not puede_validar_archivo_importacion(contexto):
        return False, "Tu rol no permite validar archivos de importación."
    if not evento_autorizado(contexto):
        return False, "El usuario no tiene acceso al evento activo."
    evento = contexto.get("evento_actual") or {}
    if evento.get("estado") != "Activo" or evento.get("fase_evento") != "Pre_evento":
        return False, "El evento debe estar Activo y en fase Pre_evento."
    return True, ""


def consultar_evento_tiene_datos(contexto: dict[str, Any], supabase: Any) -> tuple[bool | None, str]:
    """Consulta solo lectura; nunca crea, modifica ni elimina registros."""
    permitido, error = validar_contexto_importacion(contexto)
    if not permitido:
        return None, error
    evento = contexto["evento_actual"]
    cuenta_id, evento_id = int(evento["cuenta_id"]), int(evento["evento_id"])
    targets = (
        ("evp_mes_mesa", "mes_cuenta_id", "mes_evento_id", "mes_mesa_id"),
        ("evp_inv_invitacion", "inv_cuenta_id", "inv_evento_id", "inv_invitacion_id"),
        ("evp_ivt_invitado", "ivt_cuenta_id", "ivt_evento_id", "ivt_invitado_id"),
    )
    try:
        for table, account_column, event_column, selected_column in targets:
            response = (supabase.table(table).select(selected_column).eq(account_column, cuenta_id).eq(event_column, evento_id).limit(1).execute())
            if getattr(response, "data", None):
                return True, "Este evento ya contiene información y no admite importación inicial."
    except Exception as ex:
        return None, f"No fue posible comprobar si el evento está vacío ({type(ex).__name__})."
    return False, "El evento no contiene mesas, invitaciones ni invitados."


def construir_payload_importacion(preview: ImportPreview) -> dict[str, Any]:
    return {
        "version": 1,
        "mesas": [
            {"codigo_externo": table.code, "nombre": table.name}
            for table in preview.tables
        ],
        "invitaciones": [
            {
                "codigo_externo": invitation.code,
                "destinatario": invitation.recipient,
                "puestos_reservados": invitation.reserved_seats,
                "invitados": [
                    {
                        "orden": row.guest_order,
                        "nombre": row.guest_name,
                        "telefono": row.phone,
                        "email": row.email,
                        "es_principal": row.is_primary,
                        "mesa_codigo": row.table_code,
                    }
                    for row in invitation.rows
                ],
            }
            for invitation in preview.invitations
        ],
    }


def serializar_payload_importacion(payload: dict[str, Any]) -> bytes:
    return json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def calcular_hash_payload(payload: dict[str, Any]) -> str:
    return hashlib.sha256(serializar_payload_importacion(payload)).hexdigest()


def vincular_preview_contexto(preview: ImportPreview, contexto: dict[str, Any]) -> ImportPreview:
    evento = contexto.get("evento_actual") or {}
    payload = construir_payload_importacion(preview)
    return replace(
        preview,
        account_id=int(evento["cuenta_id"]),
        event_id=int(evento["evento_id"]),
        event_phase=str(evento.get("fase_evento") or ""),
        event_status=str(evento.get("estado") or ""),
        payload_hash=calcular_hash_payload(payload),
        validated_at=datetime.now(timezone.utc),
    )


def preview_coincide_contexto(preview: ImportPreview, contexto: dict[str, Any]) -> bool:
    evento = contexto.get("evento_actual") or {}
    try:
        current = (int(evento["cuenta_id"]), int(evento["evento_id"]), str(evento["fase_evento"]), str(evento["estado"]))
    except (KeyError, TypeError, ValueError):
        return False
    expected = (preview.account_id, preview.event_id, preview.event_phase, preview.event_status)
    return preview.is_valid and preview.validated_at is not None and current == expected and preview.payload_hash == calcular_hash_payload(construir_payload_importacion(preview))


_RPC_MESSAGES = {
    "IMPORT_FORBIDDEN": "No tiene permisos para importar información en este evento.",
    "EVENT_NOT_FOUND": "El evento seleccionado ya no está disponible.",
    "EVENT_NOT_ACTIVE": "La importación solo está permitida en eventos activos en Pre_evento.",
    "EVENT_NOT_PRE_EVENT": "La importación solo está permitida en eventos activos en Pre_evento.",
    "EVENT_NOT_EMPTY": "Este evento ya contiene información y no admite importación inicial.",
    "INVALID_PAYLOAD": "El archivo validado contiene información que no cumple las reglas del evento.",
    "DUPLICATE_GUEST": "El archivo validado contiene invitados duplicados.",
    "INVALID_TABLE_REFERENCE": "El archivo validado contiene una referencia de mesa inválida.",
    "CONCURRENT_IMPORT": "Otra importación fue procesada antes que esta. Actualice la información.",
    "IMPORT_INTERNAL_ERROR": "La importación no pudo completarse por un error interno. No se guardó información.",
}


def _rpc_error_code(ex: Exception) -> str:
    text = " ".join(str(getattr(ex, field, "") or "") for field in ("code", "message", "details")) + " " + str(ex)
    for code in _RPC_MESSAGES:
        if code in text:
            return code
    lowered = text.lower()
    if any(token in lowered for token in ("timeout", "network", "connection", "unreachable")):
        return "NETWORK_ERROR"
    return "RPC_ERROR"


def ejecutar_importacion(supabase: Any, contexto: dict[str, Any], preview: ImportPreview) -> ImportExecutionResult:
    context_ok, _ = validar_contexto_importacion(contexto)
    if not context_ok or not puede_ejecutar_importacion(contexto):
        return ImportExecutionResult(False, "IMPORT_FORBIDDEN", _RPC_MESSAGES["IMPORT_FORBIDDEN"])
    if not preview_coincide_contexto(preview, contexto):
        return ImportExecutionResult(False, "CONTEXT_CHANGED", "El evento cambió después de validar el archivo. Vuelva a seleccionar y validar el archivo.")
    evento = contexto["evento_actual"]
    payload = construir_payload_importacion(preview)
    payload_hash = calcular_hash_payload(payload)
    try:
        response = supabase.rpc(RPC_IMPORT_NAME, {
            "p_cuenta_id": int(evento["cuenta_id"]),
            "p_evento_id": int(evento["evento_id"]),
            "p_payload": payload,
            "p_payload_hash": payload_hash,
        }).execute()
        data = getattr(response, "data", None)
        if isinstance(data, list): data = data[0] if data else None
        if not isinstance(data, dict) or data.get("ok") is not True:
            return ImportExecutionResult(False, "RPC_ERROR", "No fue posible completar la importación.")
        return ImportExecutionResult(True, "OK", "Importación completada correctamente.", int(data.get("mesas_creadas", 0)), int(data.get("invitaciones_creadas", 0)), int(data.get("invitados_creados", 0)))
    except Exception as ex:
        code = _rpc_error_code(ex)
        print("[EXCEL_IMPORT][ERROR]", f"code={code}", f"cuenta={preview.account_id}", f"evento={preview.event_id}", f"hash={payload_hash}", f"mesas={len(preview.tables)}", f"invitaciones={len(preview.invitations)}", f"invitados={len(preview.rows)}", type(ex).__name__)
        return ImportExecutionResult(False, code, _RPC_MESSAGES.get(code, "No fue posible completar la importación." if code != "NETWORK_ERROR" else "No fue posible conectar con el servicio de importación."))


def validar_encabezados(values: tuple[Any, ...]) -> tuple[dict[str, int], list[ImportValidationError]]:
    errors: list[ImportValidationError] = []
    cleaned = [str(v or "").lstrip("\ufeff").strip() for v in values]
    normalized = [normalizar_texto(v) for v in cleaned]
    canonical_by_normalized = {normalizar_texto(h): h for h in CANONICAL_HEADERS}
    seen: set[str] = set()
    mapping: dict[str, int] = {}
    for index, key in enumerate(normalized):
        if key in seen and key:
            errors.append(ImportValidationError(1, cleaned[index], cleaned[index], "encabezado_duplicado", "El encabezado está duplicado."))
        seen.add(key)
        canonical = canonical_by_normalized.get(key)
        if canonical:
            mapping[canonical] = index
        elif key:
            errors.append(ImportValidationError(1, cleaned[index], cleaned[index], "columna_desconocida", "La columna no pertenece a la plantilla aprobada."))
    for header in CANONICAL_HEADERS:
        if header not in mapping:
            errors.append(ImportValidationError(1, header, "", "columna_faltante", f"Falta la columna esperada: {header}."))
    if len(cleaned) != len(CANONICAL_HEADERS):
        errors.append(ImportValidationError(1, "Encabezados", str(len(cleaned)), "cantidad_columnas", "Se requieren exactamente nueve columnas."))
    return mapping, errors


def _issue(row: int, column: str, value: Any, code: str, message: str) -> ImportValidationError:
    return ImportValidationError(row, column, _display(value), code, message)


def normalizar_filas(raw_rows: list[ImportRowRaw]) -> tuple[list[ImportRowNormalized], list[ImportValidationError]]:
    rows: list[ImportRowNormalized] = []
    errors: list[ImportValidationError] = []
    for raw in raw_rows:
        v = raw.values
        code_value, recipient_value = v[CANONICAL_HEADERS[0]], v[CANONICAL_HEADERS[1]]
        order_value, name_value = v[CANONICAL_HEADERS[2]], v[CANONICAL_HEADERS[3]]
        phone_value, email_value = v[CANONICAL_HEADERS[4]], v[CANONICAL_HEADERS[5]]
        primary_value, table_code_value, table_name_value = v[CANONICAL_HEADERS[6]], v[CANONICAL_HEADERS[7]], v[CANONICAL_HEADERS[8]]
        code = code_value.strip() if isinstance(code_value, str) else ""
        recipient = recipient_value.strip() if isinstance(recipient_value, str) else ""
        name = name_value.strip() if isinstance(name_value, str) else ""
        if not code: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[0], code_value, "requerido_o_tipo", "El código es obligatorio y debe ser texto."))
        if not recipient: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[1], recipient_value, "requerido_o_tipo", "El destinatario es obligatorio y debe ser texto."))
        if not name: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[3], name_value, "requerido_o_tipo", "El nombre es obligatorio y debe ser texto."))
        for column, value in ((CANONICAL_HEADERS[0], code), (CANONICAL_HEADERS[1], recipient), (CANONICAL_HEADERS[3], name)):
            if len(value) > _LIMITS[column]: errors.append(_issue(raw.excel_row, column, value, "longitud_excedida", f"Máximo {_LIMITS[column]} caracteres."))
        order = order_value if isinstance(order_value, int) and not isinstance(order_value, bool) and order_value > 0 else None
        if order is None: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[2], order_value, "orden_invalido", "El orden debe ser un entero positivo."))
        phone = phone_value.strip() if isinstance(phone_value, str) else None
        if phone_value not in (None, "") and phone is None: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[4], phone_value, "texto_requerido", "El teléfono debe estar almacenado como texto."))
        if phone and len(phone) > 20: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[4], phone, "longitud_excedida", "Máximo 20 caracteres."))
        email = email_value.strip().lower() if isinstance(email_value, str) and email_value.strip() else None
        if email_value not in (None, "") and email is None: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[5], email_value, "email_invalido", "El email debe ser texto."))
        if email and (len(email) > 254 or not _EMAIL.fullmatch(email)): errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[5], email, "email_invalido", "El formato del email no es válido."))
        primary_key = normalizar_texto(primary_value) if isinstance(primary_value, str) else ""
        primary = True if primary_key == "si" else False if primary_key == "no" else None
        if primary is None: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[6], primary_value, "principal_invalido", "Use Sí, Si o No."))
        table_code = table_code_value.strip() if isinstance(table_code_value, str) and table_code_value.strip() else None
        table_name = table_name_value.strip() if isinstance(table_name_value, str) and table_name_value.strip() else None
        if table_code_value not in (None, "") and table_code is None: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[7], table_code_value, "texto_requerido", "Mesa ID debe ser texto."))
        if table_code and not table_name: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[8], table_name_value, "nombre_mesa_requerido", "Nombre de mesa es obligatorio cuando existe Mesa ID."))
        if table_name and not table_code: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[8], table_name, "mesa_id_requerido", "Mesa ID es obligatorio cuando existe nombre de mesa."))
        if table_code and len(table_code) > 50: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[7], table_code, "longitud_excedida", "Máximo 50 caracteres."))
        if table_name and len(table_name) > 30: errors.append(_issue(raw.excel_row, CANONICAL_HEADERS[8], table_name, "longitud_excedida", "Máximo 30 caracteres."))
        rows.append(ImportRowNormalized(raw.excel_row, code, recipient, order, name, normalizar_texto(name), phone, email, primary, table_code, table_name, normalizar_texto(table_name) if table_name else None))
    return rows, errors


def validar_filas(rows: list[ImportRowNormalized]) -> list[ImportValidationError]:
    errors: list[ImportValidationError] = []
    invitations: dict[str, list[ImportRowNormalized]] = defaultdict(list)
    guest_names: dict[str, list[ImportRowNormalized]] = defaultdict(list)
    tables: dict[str, list[ImportRowNormalized]] = defaultdict(list)
    table_names: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        invitations[row.invitation_code].append(row)
        if row.guest_name_normalized: guest_names[row.guest_name_normalized].append(row)
        if row.table_code:
            tables[row.table_code].append(row)
            if row.table_name_normalized: table_names[row.table_name_normalized].add(row.table_code)
    for code, group in invitations.items():
        recipients = {normalizar_texto(r.recipient) for r in group if r.recipient}
        orders: dict[int, list[ImportRowNormalized]] = defaultdict(list)
        for row in group:
            if row.guest_order is not None: orders[row.guest_order].append(row)
        if len([r for r in group if r.is_primary is True]) != 1:
            for row in group: errors.append(_issue(row.excel_row, CANONICAL_HEADERS[6], row.is_primary, "cantidad_principales", f"La invitación {code or '(vacía)'} debe tener exactamente un principal."))
        if len(recipients) > 1:
            for row in group: errors.append(_issue(row.excel_row, CANONICAL_HEADERS[1], row.recipient, "destinatario_contradictorio", "Una invitación no puede tener destinatarios distintos."))
        for order, duplicates in orders.items():
            if len(duplicates) > 1:
                for row in duplicates: errors.append(_issue(row.excel_row, CANONICAL_HEADERS[2], order, "orden_duplicado", "El orden debe ser único dentro de la invitación."))
    for group in guest_names.values():
        if len(group) > 1:
            for row in group: errors.append(_issue(row.excel_row, CANONICAL_HEADERS[3], row.guest_name, "nombre_duplicado", "Nombre duplicado en el archivo; añada un diferenciador."))
    for code, group in tables.items():
        names = {r.table_name_normalized for r in group if r.table_name_normalized}
        if len(names) > 1:
            for row in group: errors.append(_issue(row.excel_row, CANONICAL_HEADERS[8], row.table_name, "mesa_inconsistente", "El mismo Mesa ID usa nombres diferentes."))
    for name, codes in table_names.items():
        if len(codes) > 1:
            for code in codes:
                for row in tables[code]:
                    if row.table_name_normalized == name: errors.append(_issue(row.excel_row, CANONICAL_HEADERS[8], row.table_name, "nombre_mesa_duplicado", "Dos Mesa ID distintos no pueden usar el mismo nombre."))
    return errors


def agrupar_invitaciones(rows: list[ImportRowNormalized]) -> tuple[InvitationImportGroup, ...]:
    grouped: dict[str, list[ImportRowNormalized]] = defaultdict(list)
    for row in rows: grouped[row.invitation_code].append(row)
    return tuple(InvitationImportGroup(code, group[0].recipient, tuple(group), len(group)) for code, group in grouped.items() if code)


def agrupar_mesas(rows: list[ImportRowNormalized]) -> tuple[TableImportGroup, ...]:
    grouped: dict[str, list[ImportRowNormalized]] = defaultdict(list)
    for row in rows:
        if row.table_code: grouped[row.table_code].append(row)
    return tuple(TableImportGroup(code, group[0].table_name or "", tuple(r.excel_row for r in group)) for code, group in grouped.items())


def leer_archivo_excel(filename: str, content: bytes, *, max_size: int = MAX_FILE_SIZE, max_rows: int = MAX_DATA_ROWS) -> ImportPreview:
    initial_errors: list[ImportValidationError] = []
    safe_name = Path(str(filename or "")).name
    if safe_name != filename or not safe_name.lower().endswith(".xlsx") or safe_name.lower().endswith((".xls.xlsx", ".xlsm.xlsx")):
        initial_errors.append(_issue(0, "Archivo", filename, "extension_invalida", "Solo se admiten archivos .xlsx con nombre seguro."))
    if not content: initial_errors.append(_issue(0, "Archivo", "", "archivo_vacio", "El archivo está vacío."))
    if len(content) > max_size: initial_errors.append(_issue(0, "Archivo", len(content), "tamano_excedido", f"El archivo supera {max_size} bytes."))
    if content and not is_zipfile(BytesIO(content)): initial_errors.append(_issue(0, "Archivo", safe_name, "xlsx_invalido", "El contenido no es un XLSX legible."))
    if initial_errors: return _preview(safe_name, [], [], initial_errors, [])
    workbook = formula_workbook = None
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        formula_workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if IMPORT_SHEET not in workbook.sheetnames:
            return _preview(safe_name, [], [], [_issue(1, "Hoja", IMPORT_SHEET, "hoja_faltante", f"Falta la hoja {IMPORT_SHEET}.")], [])
        sheet, formula_sheet = workbook[IMPORT_SHEET], formula_workbook[IMPORT_SHEET]
        first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        mapping, errors = validar_encabezados(first)
        if errors: return _preview(safe_name, [], [], errors, [])
        raw_rows: list[ImportRowRaw] = []
        warnings: list[ImportValidationWarning] = []
        empty_seen = False
        for excel_row, (values, formulas) in enumerate(zip(sheet.iter_rows(min_row=2, values_only=True), formula_sheet.iter_rows(min_row=2, values_only=True)), 2):
            selected = {h: values[index] if index < len(values) else None for h, index in mapping.items()}
            if all(value in (None, "") for value in selected.values()):
                empty_seen = True
                continue
            if empty_seen:
                warnings.append(ImportValidationWarning(excel_row, "Fila", "", "fila_vacia_intermedia", "Se ignoró una fila completamente vacía antes de esta fila."))
                empty_seen = False
            for header, index in mapping.items():
                formula = formulas[index] if index < len(formulas) else None
                if isinstance(formula, str) and formula.startswith("="):
                    errors.append(_issue(excel_row, header, formula, "formula_no_permitida", "No se admiten fórmulas."))
            raw_rows.append(ImportRowRaw(excel_row, selected))
            if len(raw_rows) > max_rows:
                errors.append(_issue(excel_row, "Archivo", len(raw_rows), "limite_filas", f"Se permiten como máximo {max_rows} filas."))
                break
        if not raw_rows: errors.append(_issue(0, "Archivo", safe_name, "sin_datos", "La hoja Importación no contiene filas de datos."))
        rows, row_errors = normalizar_filas(raw_rows)
        errors.extend(row_errors)
        errors.extend(validar_filas(rows))
        return _preview(safe_name, raw_rows, rows, errors, warnings)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as ex:
        return _preview(safe_name, [], [], [_issue(0, "Archivo", safe_name, "xlsx_invalido", f"No se pudo leer el XLSX ({type(ex).__name__}).")], [])
    finally:
        if formula_workbook is not None: formula_workbook.close()
        if workbook is not None: workbook.close()


def _preview(filename: str, raw: list[ImportRowRaw], rows: list[ImportRowNormalized], errors: list[ImportValidationError], warnings: list[ImportValidationWarning]) -> ImportPreview:
    invitations, tables = agrupar_invitaciones(rows), agrupar_mesas(rows)
    summary = ImportSummary(len(raw), len(invitations), len(rows), len(tables), len(errors), len(warnings))
    return ImportPreview(filename, tuple(raw), tuple(rows), invitations, tables, tuple(errors), tuple(warnings), summary)


def _csv_safe(value: Any) -> str:
    text = _display(value)
    stripped = text.lstrip()
    return "'" + text if stripped.startswith(("=", "+", "-", "@")) else text


def generar_archivo_errores(preview: ImportPreview) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(("Fila", "Columna", "Valor", "Código de error", "Mensaje", "Severidad"))
    for issue in (*preview.errors, *preview.warnings):
        writer.writerow((issue.excel_row, _csv_safe(issue.column), _csv_safe(issue.original_value), _csv_safe(issue.code), _csv_safe(issue.message), issue.severity))
    return ("\ufeff" + output.getvalue()).encode("utf-8")
